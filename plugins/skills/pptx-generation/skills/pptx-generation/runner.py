from __future__ import annotations

import io
import json
import math
import struct
import zipfile
import zlib
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as xml_escape

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.core.artifacts.store import ArtifactStore, ThreadPaths
from app.schemas import ArtifactRef

_PillowFont = ImageFont.ImageFont | ImageFont.FreeTypeFont


@dataclass
class SkillRunResult:
    skill_name: str
    outputs: list[ArtifactRef] = field(default_factory=list)
    data: dict[str, Any] = field(default_factory=dict)


@dataclass
class _LaneLayout:
    index: int
    name: str
    x: int
    y: int
    width: int
    height: int
    node_names: list[str]
    fill: str
    accent: str


@dataclass
class _NodeLayout:
    name: str
    lane_index: int
    x: int
    y: int
    width: int
    height: int
    fill: str
    accent: str


@dataclass
class _DiagramLayout:
    title: str
    width: int
    height: int
    lanes: list[_LaneLayout]
    nodes: dict[str, _NodeLayout]


_FONT_5X7 = {
    " ": ("00000", "00000", "00000", "00000", "00000", "00000", "00000"),
    "-": ("00000", "00000", "00000", "11111", "00000", "00000", "00000"),
    "/": ("00001", "00010", "00100", "01000", "10000", "00000", "00000"),
    "0": ("01110", "10001", "10011", "10101", "11001", "10001", "01110"),
    "1": ("00100", "01100", "00100", "00100", "00100", "00100", "01110"),
    "2": ("01110", "10001", "00001", "00010", "00100", "01000", "11111"),
    "3": ("11110", "00001", "00001", "01110", "00001", "00001", "11110"),
    "4": ("00010", "00110", "01010", "10010", "11111", "00010", "00010"),
    "5": ("11111", "10000", "10000", "11110", "00001", "00001", "11110"),
    "6": ("01110", "10000", "10000", "11110", "10001", "10001", "01110"),
    "7": ("11111", "00001", "00010", "00100", "01000", "01000", "01000"),
    "8": ("01110", "10001", "10001", "01110", "10001", "10001", "01110"),
    "9": ("01110", "10001", "10001", "01111", "00001", "00001", "01110"),
    "A": ("01110", "10001", "10001", "11111", "10001", "10001", "10001"),
    "B": ("11110", "10001", "10001", "11110", "10001", "10001", "11110"),
    "C": ("01110", "10001", "10000", "10000", "10000", "10001", "01110"),
    "D": ("11110", "10001", "10001", "10001", "10001", "10001", "11110"),
    "E": ("11111", "10000", "10000", "11110", "10000", "10000", "11111"),
    "F": ("11111", "10000", "10000", "11110", "10000", "10000", "10000"),
    "G": ("01110", "10001", "10000", "10111", "10001", "10001", "01110"),
    "H": ("10001", "10001", "10001", "11111", "10001", "10001", "10001"),
    "I": ("01110", "00100", "00100", "00100", "00100", "00100", "01110"),
    "J": ("00111", "00010", "00010", "00010", "10010", "10010", "01100"),
    "K": ("10001", "10010", "10100", "11000", "10100", "10010", "10001"),
    "L": ("10000", "10000", "10000", "10000", "10000", "10000", "11111"),
    "M": ("10001", "11011", "10101", "10101", "10001", "10001", "10001"),
    "N": ("10001", "11001", "10101", "10011", "10001", "10001", "10001"),
    "O": ("01110", "10001", "10001", "10001", "10001", "10001", "01110"),
    "P": ("11110", "10001", "10001", "11110", "10000", "10000", "10000"),
    "Q": ("01110", "10001", "10001", "10001", "10101", "10010", "01101"),
    "R": ("11110", "10001", "10001", "11110", "10100", "10010", "10001"),
    "S": ("01111", "10000", "10000", "01110", "00001", "00001", "11110"),
    "T": ("11111", "00100", "00100", "00100", "00100", "00100", "00100"),
    "U": ("10001", "10001", "10001", "10001", "10001", "10001", "01110"),
    "V": ("10001", "10001", "10001", "10001", "10001", "01010", "00100"),
    "W": ("10001", "10001", "10001", "10101", "10101", "10101", "01010"),
    "X": ("10001", "10001", "01010", "00100", "01010", "10001", "10001"),
    "Y": ("10001", "10001", "01010", "00100", "00100", "00100", "00100"),
    "Z": ("11111", "00001", "00010", "00100", "01000", "10000", "11111"),
}


class SkillRunner:
    """Deterministic local skill runner with content-rich artifact outputs."""

    def __init__(self, artifact_store: ArtifactStore) -> None:
        self.artifact_store = artifact_store

    def run(self, skill_name: str, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        if skill_name == "drawio-generation":
            return self._drawio(spec, paths)
        if skill_name == "markdown-rendering":
            return self._markdown(spec, paths)
        if skill_name == "excel-generation":
            return self._excel(spec, paths)
        if skill_name == "pptx-generation":
            return self._pptx(spec, paths)
        if skill_name == "xmind-generation":
            return self._xmind(spec, paths)
        if skill_name == "behavior-detection":
            return self._behavior(spec, paths)
        if skill_name == "behavior-review":
            return self._behavior_review(spec, paths)
        raise KeyError(f"Unsupported skill: {skill_name}")

    def _drawio(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        title = str(spec.get("title") or "JetLinks Agent Runtime v2")
        nodes = [str(item) for item in spec.get("nodes", [])] or [
            "用户入口",
            "无状态 Agent Runtime",
            "Skill Runner",
            "Artifact Store",
            "内容级 Verifier",
        ]
        edges = self._string_pairs(spec.get("edges")) or [
            ["用户入口", "无状态 Agent Runtime"],
            ["无状态 Agent Runtime", "Skill Runner"],
            ["Skill Runner", "Artifact Store"],
            ["Artifact Store", "内容级 Verifier"],
        ]
        swimlanes = [str(item) for item in spec.get("swimlanes", [])] or ["交互层", "Agent 编排层", "技能执行层", "质量层"]
        polished = spec.get("visual_style") == "polished" or bool(spec.get("refinement_requested"))
        colors = dict(spec.get("color_semantics", {}))
        palette = [
            str(colors.get("entry", "#dbeafe")),
            str(colors.get("agent", "#dcfce7")),
            str(colors.get("skill", "#fef3c7")),
            str(colors.get("artifact", "#fce7f3")),
            str(colors.get("quality", "#fee2e2")),
        ]
        lane_node_hints = self._lane_node_hints(spec.get("lane_nodes"), nodes, swimlanes)
        layout = self._drawio_layout(title, nodes, edges, swimlanes, palette, lane_node_hints)

        cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
        lane_ids: list[str] = []
        for lane in layout.lanes:
            lane_id = f"lane{lane.index + 1}"
            lane_ids.append(lane_id)
            lane_style = (
                "swimlane;whiteSpace=wrap;html=1;startSize=46;rounded=1;arcSize=8;"
                "fontStyle=1;fontSize=15;fontFamily=PingFang SC;spacing=12;"
                f"fillColor={lane.fill};strokeColor=#cbd5e1;fontColor=#0f172a;"
                "shadow=1;swimlaneFillColor=#ffffff;"
                if polished
                else (
                    "swimlane;whiteSpace=wrap;html=1;startSize=42;rounded=1;arcSize=6;"
                    "fontStyle=1;fontSize=14;fontFamily=PingFang SC;spacing=10;"
                    f"fillColor={lane.fill};strokeColor=#94a3b8;fontColor=#0f172a;"
                )
            )
            cells.append(
                f'<mxCell id="{lane_id}" value="{xml_escape(lane.name)}" '
                f'style="{lane_style}" vertex="1" parent="1">'
                f'<mxGeometry x="{lane.x}" y="{lane.y}" width="{lane.width}" height="{lane.height}" '
                'as="geometry"/></mxCell>'
            )

        node_ids: dict[str, str] = {}
        for idx, node in enumerate(nodes):
            node_layout = layout.nodes[node]
            lane = layout.lanes[node_layout.lane_index]
            node_id = f"n{idx + 1}"
            node_ids[node] = node_id
            node_style = (
                "rounded=1;whiteSpace=wrap;html=1;arcSize=10;spacing=10;fontStyle=1;fontSize=13;"
                "fontFamily=PingFang SC;"
                f"fillColor={node_layout.fill};strokeColor={node_layout.accent};fontColor=#0f172a;shadow=1;"
                if polished
                else (
                    "rounded=1;whiteSpace=wrap;html=1;arcSize=8;spacing=8;fontSize=12;"
                    "fontFamily=PingFang SC;"
                    f"fillColor={node_layout.fill};strokeColor={node_layout.accent};fontColor=#0f172a;"
                )
            )
            cells.append(
                f'<mxCell id="{node_id}" value="{xml_escape(node)}" '
                f'style="{node_style}" vertex="1" parent="{lane_ids[node_layout.lane_index]}">'
                f'<mxGeometry x="{node_layout.x - lane.x}" y="{node_layout.y - lane.y}" '
                f'width="{node_layout.width}" height="{node_layout.height}" as="geometry"/></mxCell>'
            )

        for idx, edge in enumerate(edges, start=1):
            source = node_ids.get(edge[0], "n1")
            target = node_ids.get(edge[1], "n2")
            edge_style = (
                "edgeStyle=orthogonalEdgeStyle;rounded=1;orthogonalLoop=1;jettySize=auto;html=1;"
                "endArrow=block;endFill=1;strokeWidth=2;strokeColor=#2563eb;"
                "exitX=1;exitY=0.5;exitDx=0;exitDy=0;entryX=0;entryY=0.5;entryDx=0;entryDy=0;"
                if polished
                else (
                    "edgeStyle=orthogonalEdgeStyle;rounded=1;orthogonalLoop=1;jettySize=auto;html=1;"
                    "endArrow=block;endFill=1;strokeWidth=2;strokeColor=#64748b;"
                )
            )
            cells.append(
                f'<mxCell id="e{idx}" value="" style="{edge_style}" edge="1" parent="1" '
                f'source="{source}" target="{target}"><mxGeometry relative="1" as="geometry"/></mxCell>'
            )

        xml = (
            '<mxfile host="JetLinks Agent Runtime v2">'
            f'<diagram name="{xml_escape(title)}"><mxGraphModel grid="1" gridSize="10" page="1" '
            f'pageWidth="{layout.width}" pageHeight="{layout.height}">'
            f"<root>{''.join(cells)}</root></mxGraphModel></diagram></mxfile>"
        )
        base_name = self._next_drawio_base_name(
            paths,
            "prototype" if spec.get("diagram_type") == "prototype_wireframe" else "architecture",
        )
        filename = f"{base_name}.drawio"
        artifact = self.artifact_store.write_text_artifact(paths, filename, xml)
        png_filename = filename.removesuffix(".drawio") + ".png"
        png_artifact = self.artifact_store.write_bytes_artifact(
            paths,
            png_filename,
            self._drawio_preview_png(layout, edges, polished=polished),
        )
        return SkillRunResult(
            "drawio-generation",
            [artifact, png_artifact],
            {
                "nodes": nodes,
                "edges": edges,
                "swimlanes": swimlanes,
                "lane_nodes": lane_node_hints,
                "color_semantics": colors,
                "preview": png_artifact.name,
            },
        )

    def _markdown(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        title = str(spec.get("title") or "JetLinks Agent Runtime v2")
        summary = str(spec.get("summary") or "基于用户需求生成可验证的 Markdown 产物。")
        body = f"""# {title}

## 目标

{summary}

## 架构

```mermaid
flowchart LR
  U[用户自然语言] --> S[结构化 Spec]
  S --> K[Skill 执行]
  K --> A[Artifact 输出]
  A --> P[文件预览]
  A --> V[内容级 Verifier]
  V -->|合格| R[Agent 回复]
  V -->|不合格| T[自动修正重试]
  T --> K
```

## 接口契约

| 字段 | 类型 | 说明 |
|---|---|---|
| agent | string | 当前 Agent 配置名称 |
| thread_id | string | 文件容器标识，不作为长期记忆 |
| artifacts | array | 生成文件、预览地址和下载地址 |
| verification | object | 内容级校验结果 |

## ECharts 配置块

```json echarts
{{
  "xAxis": {{"type": "category", "data": ["Spec", "Skill", "Artifact", "Verify"]}},
  "yAxis": {{"type": "value"}},
  "series": [{{"type": "bar", "data": [1, 1, 1, 1]}}]
}}
```

## 附件块

> artifact://runtime/generated
"""
        artifact = self.artifact_store.write_text_artifact(paths, "result.md", body)
        return SkillRunResult("markdown-rendering", [artifact], {"format": "markdown", "sections": 5})

    def _excel(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        fields = list(spec.get("fields", []))
        rows = [
            ["设备ID", "设备名称", "产品类型", "在线状态", "最后活跃时间"],
            ["dev-001", "一号网关", "网关", "在线", "2026-04-29 10:00:00"],
            ["dev-002", "二号传感器", "传感器", "离线", "2026-04-28 18:30:00"],
            ["dev-003", "围栏摄像头", "视频设备", "告警", "2026-04-29 11:20:00"],
        ]
        summary_rows = [
            ["指标", "公式"],
            ["设备总数", "=COUNTA('设备模板'!A2:A100)"],
            ["在线数量", '=COUNTIF(\'设备模板\'!D2:D100,"在线")'],
            ["告警数量", '=COUNTIF(\'设备模板\'!D2:D100,"告警")'],
        ]
        artifact = self.artifact_store.write_bytes_artifact(
            paths,
            "device-template.xlsx",
            self._build_xlsx(rows, summary_rows),
        )
        return SkillRunResult(
            "excel-generation",
            [artifact],
            {
                "sheets": ["设备模板", "状态汇总"],
                "fields": fields,
                "formulas": ["COUNTA", "COUNTIF"],
                "freeze_panes": "A2",
                "validation": "在线状态: 在线/离线/告警",
            },
        )

    def _pptx(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        slides = list(spec.get("slides", [])) or [
            {"title": "JetLinks Agent Runtime v2", "kind": "cover"},
            {"title": "目录", "kind": "agenda"},
            {"title": "总体架构", "kind": "architecture"},
            {"title": "生成链路", "kind": "process"},
            {"title": "实施计划", "kind": "plan"},
            {"title": "风险与控制", "kind": "risks"},
            {"title": "总结", "kind": "summary"},
        ]
        normalized: list[dict[str, Any]] = []
        for idx, slide in enumerate(slides, start=1):
            if not isinstance(slide, dict):
                continue
            kind = str(slide.get("kind", "content"))
            normalized.append(
                {
                    "title": str(slide.get("title", f"Slide {idx}")),
                    "kind": kind,
                    "bullets": self._pptx_bullets(kind, slide.get("bullets")),
                }
            )
        artifact = self.artifact_store.write_bytes_artifact(paths, "deck.pptx", self._build_pptx(normalized))
        return SkillRunResult("pptx-generation", [artifact], {"slides": len(normalized), "slide_kinds": normalized})

    def _xmind(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        root_topic = str(spec.get("root_topic") or spec.get("title") or "JetLinks Agent Runtime v2")
        topics = list(spec.get("topics", []))
        content = [
            {
                "id": "sheet-1",
                "class": "sheet",
                "title": root_topic,
                "rootTopic": {
                    "id": "root",
                    "class": "topic",
                    "title": root_topic,
                    "children": {
                        "attached": [
                            {
                                "id": f"topic-{idx}",
                                "class": "topic",
                                "title": str(topic.get("title", f"主题 {idx}")),
                                "notes": {"plain": {"content": "由 JetLinks Agent Runtime v2 生成。"}},
                                "children": {
                                    "attached": [
                                        {
                                            "id": f"topic-{idx}-{child_idx}",
                                            "class": "topic",
                                            "title": str(child),
                                            "markers": ["priority-1"] if child_idx == 1 else [],
                                        }
                                        for child_idx, child in enumerate(topic.get("children", []), start=1)
                                    ]
                                },
                            }
                            for idx, topic in enumerate(topics, start=1)
                            if isinstance(topic, dict)
                        ]
                    },
                },
            }
        ]
        metadata = {"creator": {"name": "JetLinks Agent Runtime v2"}, "activeSheetId": "sheet-1"}
        artifact = self.artifact_store.write_bytes_artifact(
            paths,
            "mindmap.xmind",
            self._zip_bytes({"content.json": json.dumps(content, ensure_ascii=False), "metadata.json": json.dumps(metadata)}),
        )
        return SkillRunResult("xmind-generation", [artifact], {"root_topic": root_topic, "topics": len(topics)})

    def _behavior(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        has_visual_evidence = bool(spec.get("has_visual_evidence"))
        candidates = [str(item) for item in spec.get("text_rule_candidates", [])] or ["待确认行为类型"]
        if not has_visual_evidence:
            data = {
                "decision": "needs_visual_confirmation",
                "category": candidates[0],
                "confidence": 0.0,
                "text_rule_confidence": 0.72,
                "evidence_mode": "text_only",
                "reason": "仅收到文本描述，不能输出视觉识别结果，需要上传图片/视频或结构化视觉证据确认。",
            }
            return self._behavior_detection_result(paths, data)
        data = {
            "decision": "critical",
            "category": candidates[0],
            "confidence": 0.88,
            "evidence_mode": "visual_or_structured",
            "reason": "已提供视觉或结构化证据，可以进行行为识别判断。",
        }
        return self._behavior_detection_result(paths, data)

    def _behavior_detection_result(self, paths: ThreadPaths, payload: dict[str, Any]) -> SkillRunResult:
        md_artifact = self.artifact_store.write_text_artifact(paths, "behavior-detection.md", self._behavior_detection_markdown(payload))
        json_artifact = self.artifact_store.write_text_artifact(
            paths,
            "behavior-detection.json",
            json.dumps(payload, ensure_ascii=False, indent=2),
        )
        return SkillRunResult("behavior-detection", [md_artifact, json_artifact], payload)

    @staticmethod
    def _behavior_detection_markdown(payload: dict[str, Any]) -> str:
        return "\n".join(
            [
                "# 行为识别报告",
                "",
                "## 结论",
                "",
                f"- 决策：{payload.get('decision')}",
                f"- 类别：{payload.get('category')}",
                f"- 置信度：{payload.get('confidence')}",
                f"- 证据模式：{payload.get('evidence_mode')}",
                "",
                "## 说明",
                "",
                str(payload.get("reason") or ""),
                "",
            ]
        )

    def _behavior_review(self, spec: dict[str, Any], paths: ThreadPaths) -> SkillRunResult:
        has_visual_evidence = bool(spec.get("has_visual_evidence"))
        candidates = [str(item) for item in spec.get("text_rule_candidates", [])] or ["待确认行为类型"]
        event_id = str(spec.get("event_id") or "待生成")
        original_decision = str(spec.get("original_decision") or ("critical" if has_visual_evidence else "needs_visual_confirmation"))
        payload = {
            "event_id": event_id,
            "category": candidates[0],
            "original_decision": original_decision,
            "review_decision": "confirm_incident" if has_visual_evidence else "need_more_evidence",
            "risk_level": "high" if has_visual_evidence else "pending",
            "evidence_mode": "visual_or_structured" if has_visual_evidence else "text_only",
            "confidence": 0.86 if has_visual_evidence else 0.0,
            "text_rule_confidence": 0.72,
            "evidence_gaps": [] if has_visual_evidence else ["缺少图片、视频或结构化视觉证据", "不能仅凭文本规则输出视觉置信度"],
            "actions": (
                ["生成告警工单", "保留证据截图/视频片段", "通知值班人员复核", "记录处置闭环"]
                if has_visual_evidence
                else ["请求上传图片/视频/结构化检测结果", "保留文本规则命中记录", "暂不升级为视觉确认事件"]
            ),
            "manual_review_required": not has_visual_evidence,
            "reason": "复判遵循证据优先原则：有视觉证据才确认事件，缺证据时只给文本规则命中和补证建议。",
        }
        md_artifact = self.artifact_store.write_text_artifact(paths, "behavior-review.md", self._behavior_review_markdown(payload))
        json_artifact = self.artifact_store.write_text_artifact(
            paths,
            "behavior-review.json",
            json.dumps(payload, ensure_ascii=False, indent=2),
        )
        return SkillRunResult("behavior-review", [md_artifact, json_artifact], payload)

    @staticmethod
    def _behavior_review_markdown(payload: dict[str, Any]) -> str:
        lines = [
            "# 行为识别复判报告",
            "",
            "## 结论",
            "",
            f"- 事件ID: {payload['event_id']}",
            f"- 行为类型: {payload['category']}",
            f"- 原始判断: {payload['original_decision']}",
            f"- 复判结论: {payload['review_decision']}",
            f"- 风险等级: {payload['risk_level']}",
            f"- 证据模式: {payload['evidence_mode']}",
            f"- 视觉置信度: {payload['confidence']}",
            f"- 文本规则置信度: {payload['text_rule_confidence']}",
            "",
            "## 证据缺口",
            "",
        ]
        for item in payload["evidence_gaps"] or ["暂无缺口"]:
            lines.append(f"- {item}")
        lines.extend(["", "## 处置建议", ""])
        for item in payload["actions"]:
            lines.append(f"- {item}")
        lines.extend(
            [
                "",
                "## 复判规则",
                "",
                "- 没有图片、视频或结构化视觉证据时，不输出视觉确认结论。",
                "- 文本规则命中只能作为疑似告警，不能替代视觉证据。",
                "- 高风险事件需要保留证据、处理动作和复判人/复判系统记录。",
            ]
        )
        return "\n".join(lines) + "\n"

    @staticmethod
    def _string_pairs(value: object) -> list[list[str]]:
        if not isinstance(value, list):
            return []
        pairs: list[list[str]] = []
        for item in value:
            if isinstance(item, list | tuple) and len(item) >= 2:
                pairs.append([str(item[0]), str(item[1])])
        return pairs

    @staticmethod
    def _zip_bytes(files: Mapping[str, str | bytes]) -> bytes:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for path, content in files.items():
                archive.writestr(path, content)
        return buffer.getvalue()

    @staticmethod
    def _next_drawio_base_name(paths: ThreadPaths, stem: str) -> str:
        if not (paths.outputs / f"{stem}.drawio").exists() and not (paths.outputs / f"{stem}.png").exists():
            return stem
        version = 2
        while (paths.outputs / f"{stem}_v{version}.drawio").exists() or (
            paths.outputs / f"{stem}_v{version}.png"
        ).exists():
            version += 1
        return f"{stem}_v{version}"

    @classmethod
    def _drawio_layout(
        cls,
        title: str,
        nodes: list[str],
        edges: list[list[str]],
        swimlanes: list[str],
        palette: list[str],
        lane_node_hints: dict[str, list[str]] | None = None,
    ) -> _DiagramLayout:
        lane_names = swimlanes or ["架构层"]
        lane_count = max(1, len(lane_names))
        assignments = cls._assign_nodes_to_lanes(nodes, edges, lane_names, lane_node_hints)
        lane_nodes: list[list[str]] = [[] for _ in lane_names]
        hinted_nodes: set[str] = set()
        if lane_node_hints:
            for lane_idx, lane_name in enumerate(lane_names):
                for node in lane_node_hints.get(lane_name, []):
                    if node in hinted_nodes:
                        continue
                    lane_nodes[lane_idx].append(node)
                    hinted_nodes.add(node)
        for idx, node in enumerate(nodes):
            if node in hinted_nodes:
                continue
            lane_idx = assignments.get(node, cls._fallback_lane_index(idx, len(nodes), lane_count))
            lane_nodes[lane_idx].append(node)

        lane_width = 292 if lane_count <= 4 else 268
        lane_gap = 28 if lane_count <= 4 else 22
        node_width = lane_width - 54
        node_height = 70
        row_gap = 22
        header_height = 46
        margin_x = 48
        margin_top = 42
        title_height = 58
        footer_height = 34
        max_rows = max(1, max(len(items) for items in lane_nodes))
        lane_height = header_height + 28 + max_rows * node_height + (max_rows - 1) * row_gap + 30
        width = margin_x * 2 + lane_count * lane_width + (lane_count - 1) * lane_gap
        height = margin_top + title_height + lane_height + footer_height
        lane_y = margin_top + title_height

        lanes: list[_LaneLayout] = []
        node_layouts: dict[str, _NodeLayout] = {}
        for lane_idx, lane_name in enumerate(lane_names):
            lane_x = margin_x + lane_idx * (lane_width + lane_gap)
            accent = cls._accent_color(lane_idx)
            lane_fill = cls._lane_fill_color(lane_idx)
            lanes.append(
                _LaneLayout(
                    index=lane_idx,
                    name=lane_name,
                    x=lane_x,
                    y=lane_y,
                    width=lane_width,
                    height=lane_height,
                    node_names=lane_nodes[lane_idx],
                    fill=lane_fill,
                    accent=accent,
                )
            )
            for row, node in enumerate(lane_nodes[lane_idx]):
                node_x = lane_x + (lane_width - node_width) // 2
                node_y = lane_y + header_height + 28 + row * (node_height + row_gap)
                node_layouts[node] = _NodeLayout(
                    name=node,
                    lane_index=lane_idx,
                    x=node_x,
                    y=node_y,
                    width=node_width,
                    height=node_height,
                    fill=palette[min(lane_idx, len(palette) - 1)],
                    accent=accent,
                )

        return _DiagramLayout(title=title, width=width, height=height, lanes=lanes, nodes=node_layouts)

    @classmethod
    def _assign_nodes_to_lanes(
        cls,
        nodes: list[str],
        edges: list[list[str]],
        lane_names: list[str],
        lane_node_hints: dict[str, list[str]] | None = None,
    ) -> dict[str, int]:
        lane_count = max(1, len(lane_names))
        lane_keywords = [cls._lane_keywords(lane) for lane in lane_names]
        incoming: dict[str, int] = {node: 0 for node in nodes}
        outgoing: dict[str, int] = {node: 0 for node in nodes}
        for edge in edges:
            if len(edge) < 2:
                continue
            outgoing[edge[0]] = outgoing.get(edge[0], 0) + 1
            incoming[edge[1]] = incoming.get(edge[1], 0) + 1

        assignments: dict[str, int] = {}
        if lane_node_hints:
            for lane_idx, lane_name in enumerate(lane_names):
                for node in lane_node_hints.get(lane_name, []):
                    if node in incoming:
                        assignments[node] = lane_idx
        for idx, node in enumerate(nodes):
            if node in assignments:
                continue
            scores = [
                cls._node_lane_score(node, keywords, lane_name)
                for keywords, lane_name in zip(lane_keywords, lane_names, strict=True)
            ]
            best_score = max(scores) if scores else 0
            if best_score > 0:
                assignments[node] = scores.index(best_score)
                continue
            assignments[node] = cls._fallback_lane_index(idx, len(nodes), lane_count)

        # Source-like unknown nodes should stay near the left; sink-like unknown nodes can move right.
        for node, lane_idx in list(assignments.items()):
            if cls._node_lane_score(node, lane_keywords[lane_idx], lane_names[lane_idx]) > 0:
                continue
            if incoming.get(node, 0) == 0 and outgoing.get(node, 0) > 0:
                assignments[node] = min(assignments[node], 0)
            elif outgoing.get(node, 0) == 0 and incoming.get(node, 0) > 0:
                assignments[node] = max(assignments[node], lane_count - 1)
        return assignments

    @classmethod
    def _lane_keywords(cls, lane_name: str) -> tuple[str, ...]:
        text = lane_name.lower()
        groups: list[tuple[tuple[str, ...], tuple[str, ...]]] = [
            (
                ("接入", "入口", "设备", "网关", "协议", "连接", "采集", "api", "mqtt", "http", "coap", "modbus"),
                ("接入", "入口", "设备", "网关", "协议", "连接", "采集", "api", "mqtt", "http", "coap", "modbus", "opc"),
            ),
            (
                ("规则", "引擎", "计算", "处理", "编排", "核心", "agent", "skill", "runtime"),
                ("规则", "引擎", "计算", "处理", "过滤", "脚本", "转发", "编排", "核心", "agent", "builder", "registry", "runner", "skill", "runtime"),
            ),
            (
                ("数据", "存储", "数据库", "缓存", "时序", "服务", "artifact", "preview"),
                ("数据", "存储", "数据库", "缓存", "时序", "关系", "postgres", "postgresql", "redis", "influx", "clickhouse", "mysql", "artifact", "preview", "文件"),
            ),
            (
                ("告警", "监控", "运维", "通知", "质量", "校验", "闭环", "看板"),
                ("告警", "监控", "运维", "通知", "prometheus", "grafana", "日志", "审计", "可视化", "看板", "质量", "校验", "重试", "reply", "最终回复"),
            ),
            (
                ("交互", "页面", "前端", "用户", "输出", "状态"),
                ("交互", "页面", "前端", "用户", "工作台", "聊天", "输入", "输出", "状态", "导航", "下载", "错误"),
            ),
            (
                ("应用", "业务", "控制台", "管理", "看板"),
                ("应用", "业务", "控制台", "管理", "看板", "前端", "用户", "报表", "运营", "运维", "监控"),
            ),
        ]
        keywords: list[str] = []
        for markers, terms in groups:
            if any(marker in text for marker in markers):
                keywords.extend(terms)
        keywords.extend(cls._label_terms(lane_name))
        return tuple(dict.fromkeys(term for term in keywords if term))

    @staticmethod
    def _label_terms(label: str) -> list[str]:
        normalized = label.lower()
        for separator in ("-", "_", "/", "|", "、", "与", "和", "层", "服务"):
            normalized = normalized.replace(separator, " ")
        terms = [term.strip() for term in normalized.split() if term.strip()]
        compact = label.lower().replace(" ", "")
        if len(compact) >= 2:
            terms.append(compact)
        return [term for term in terms if len(term) >= 2 or term.isascii()]

    @staticmethod
    def _node_lane_score(node: str, keywords: tuple[str, ...], lane_name: str) -> int:
        text = node.lower()
        lane_text = lane_name.lower()
        score = 0
        for keyword in keywords:
            if keyword and keyword in text:
                score += 2 if len(keyword) <= 3 else 3
        for marker in ("接入", "规则", "数据", "存储", "告警", "监控", "运维", "质量", "交互"):
            if marker in lane_text and marker in text:
                score += 4
        if "告警" in text and any(marker in lane_text for marker in ("告警", "监控", "运维", "闭环")):
            score += 8
        if any(marker in text for marker in ("控制台", "管理后台", "管理页面")) and any(
            marker in lane_text for marker in ("应用", "业务", "监控", "运维")
        ):
            score += 8
        if any(marker in text for marker in ("数据转发", "消息转发", "协议解析", "解析器")) and any(
            marker in lane_text for marker in ("规则", "处理", "核心", "引擎")
        ):
            score += 8
        if "设备影子" in text and any(marker in lane_text for marker in ("数据", "存储", "核心", "处理")):
            score += 8
        return score

    @classmethod
    def _lane_node_hints(
        cls,
        value: object,
        nodes: list[str],
        swimlanes: list[str],
    ) -> dict[str, list[str]]:
        if not swimlanes:
            return {}
        node_set = set(nodes)
        raw_items: list[tuple[str, object]] = []
        if isinstance(value, dict):
            raw_items.extend((str(lane), raw_nodes) for lane, raw_nodes in value.items())
        elif isinstance(value, list):
            for item in value:
                if not isinstance(item, dict):
                    continue
                lane = str(item.get("lane") or item.get("swimlane") or item.get("name") or "").strip()
                raw_nodes = item.get("nodes") or item.get("items")
                raw_items.append((lane, raw_nodes))

        hints: dict[str, list[str]] = {lane: [] for lane in swimlanes}
        normalized_lanes = {cls._normalize_label(lane): lane for lane in swimlanes}
        for raw_lane, raw_nodes in raw_items:
            matched_lane = normalized_lanes.get(cls._normalize_label(raw_lane))
            if matched_lane is None or not isinstance(raw_nodes, list):
                continue
            for node in raw_nodes:
                node_name = str(node).strip()
                if node_name in node_set and node_name not in hints[matched_lane]:
                    hints[matched_lane].append(node_name)
        return {lane: lane_nodes for lane, lane_nodes in hints.items() if lane_nodes}

    @staticmethod
    def _normalize_label(value: str) -> str:
        return value.lower().replace(" ", "").replace("-", "").replace("_", "")

    @staticmethod
    def _fallback_lane_index(index: int, total: int, lane_count: int) -> int:
        if lane_count <= 1 or total <= 1:
            return 0
        return min(lane_count - 1, round(index * (lane_count - 1) / (total - 1)))

    @staticmethod
    def _accent_color(index: int) -> str:
        accents = ["#2563eb", "#16a34a", "#d97706", "#7c3aed", "#dc2626", "#0891b2"]
        return accents[index % len(accents)]

    @staticmethod
    def _lane_fill_color(index: int) -> str:
        fills = ["#eff6ff", "#f0fdf4", "#fffbeb", "#f5f3ff", "#fff1f2", "#ecfeff"]
        return fills[index % len(fills)]

    @classmethod
    def _drawio_preview_png(cls, layout: _DiagramLayout, edges: list[list[str]], polished: bool = False) -> bytes:
        scale = 2
        image = Image.new("RGB", (layout.width * scale, layout.height * scale), "#f6f8fb" if polished else "#f8fafc")
        draw = ImageDraw.Draw(image)
        title_font = cls._load_font(24 * scale, bold=True)
        lane_font = cls._load_font(16 * scale, bold=True)
        node_font = cls._load_font(14 * scale, bold=True)
        meta_font = cls._load_font(11 * scale)
        edge_color = "#2563eb" if polished else "#64748b"

        cls._draw_title(draw, layout, title_font, meta_font, scale)
        for lane in layout.lanes:
            cls._draw_lane(draw, lane, lane_font, meta_font, scale, polished=polished)
        for edge in edges:
            if len(edge) < 2:
                continue
            source = layout.nodes.get(edge[0])
            target = layout.nodes.get(edge[1])
            if source is None or target is None:
                continue
            cls._draw_edge(draw, source, target, edge_color, scale)
        for node in layout.nodes.values():
            cls._draw_node(draw, node, node_font, scale, polished=polished)

        footer = "JetLinks Agent Runtime · Draw.io source + PNG preview"
        draw.text((48 * scale, (layout.height - 24) * scale), footer, font=meta_font, fill="#64748b")
        image = image.resize((layout.width, layout.height), Image.Resampling.LANCZOS)
        buffer = io.BytesIO()
        image.save(buffer, format="PNG", optimize=True)
        return buffer.getvalue()

    @classmethod
    def _draw_title(
        cls,
        draw: ImageDraw.ImageDraw,
        layout: _DiagramLayout,
        title_font: _PillowFont,
        meta_font: _PillowFont,
        scale: int,
    ) -> None:
        title = cls._short_title(layout.title)
        draw.text((48 * scale, 30 * scale), title, font=title_font, fill="#0f172a")
        subtitle = f"{len(layout.lanes)} 个分层 · {len(layout.nodes)} 个节点 · 正交连线"
        draw.text((48 * scale, 61 * scale), subtitle, font=meta_font, fill="#64748b")

    @staticmethod
    def _short_title(title: str) -> str:
        compact = " ".join(title.split())
        if len(compact) <= 36:
            return compact or "JetLinks 架构图"
        return f"{compact[:35]}..."

    @classmethod
    def _draw_lane(
        cls,
        draw: ImageDraw.ImageDraw,
        lane: _LaneLayout,
        lane_font: _PillowFont,
        meta_font: _PillowFont,
        scale: int,
        polished: bool,
    ) -> None:
        box = cls._scale_box((lane.x, lane.y, lane.x + lane.width, lane.y + lane.height), scale)
        shadow = cls._scale_box((lane.x + 5, lane.y + 7, lane.x + lane.width + 5, lane.y + lane.height + 7), scale)
        if polished:
            draw.rounded_rectangle(shadow, radius=18 * scale, fill="#d9e2ec")
        draw.rounded_rectangle(box, radius=18 * scale, fill="#ffffff", outline="#d7dee8", width=1 * scale)
        header = cls._scale_box((lane.x, lane.y, lane.x + lane.width, lane.y + 46), scale)
        draw.rounded_rectangle(header, radius=18 * scale, fill=lane.fill)
        draw.rectangle(
            cls._scale_box((lane.x, lane.y + 28, lane.x + lane.width, lane.y + 46), scale),
            fill=lane.fill,
        )
        draw.rounded_rectangle(
            cls._scale_box((lane.x + 14, lane.y + 14, lane.x + 24, lane.y + 24), scale),
            radius=5 * scale,
            fill=lane.accent,
        )
        draw.text((lane.x * scale + 34 * scale, lane.y * scale + 11 * scale), lane.name, font=lane_font, fill="#0f172a")
        count_text = f"{len(lane.node_names)} nodes"
        count_width = int(draw.textlength(count_text, font=meta_font))
        draw.text(
            ((lane.x + lane.width) * scale - count_width - 16 * scale, lane.y * scale + 15 * scale),
            count_text,
            font=meta_font,
            fill="#64748b",
        )

    @classmethod
    def _draw_node(
        cls,
        draw: ImageDraw.ImageDraw,
        node: _NodeLayout,
        node_font: _PillowFont,
        scale: int,
        polished: bool,
    ) -> None:
        if polished:
            shadow = cls._scale_box((node.x + 4, node.y + 6, node.x + node.width + 4, node.y + node.height + 6), scale)
            draw.rounded_rectangle(shadow, radius=14 * scale, fill="#dbe3ee")
        box = cls._scale_box((node.x, node.y, node.x + node.width, node.y + node.height), scale)
        draw.rounded_rectangle(box, radius=14 * scale, fill=node.fill, outline=node.accent, width=2 * scale)
        accent = cls._scale_box((node.x, node.y, node.x + 7, node.y + node.height), scale)
        draw.rounded_rectangle(accent, radius=7 * scale, fill=node.accent)
        draw.rectangle(cls._scale_box((node.x + 4, node.y, node.x + 9, node.y + node.height), scale), fill=node.accent)
        lines = cls._wrap_text(draw, node.name, node_font, (node.width - 32) * scale, max_lines=2)
        line_height = cls._line_height(node_font) + 3 * scale
        total_height = len(lines) * line_height - 3 * scale
        text_y = node.y * scale + (node.height * scale - total_height) // 2
        for line in lines:
            line_width = int(draw.textlength(line, font=node_font))
            text_x = node.x * scale + 18 * scale + ((node.width - 28) * scale - line_width) // 2
            draw.text((text_x, text_y), line, font=node_font, fill="#0f172a")
            text_y += line_height

    @classmethod
    def _draw_edge(
        cls,
        draw: ImageDraw.ImageDraw,
        source: _NodeLayout,
        target: _NodeLayout,
        color: str,
        scale: int,
    ) -> None:
        points = cls._edge_points(source, target)
        scaled_points = [(x * scale, y * scale) for x, y in points]
        draw.line(scaled_points, fill="#d8e0ea", width=5 * scale, joint="curve")
        draw.line(scaled_points, fill=color, width=2 * scale, joint="curve")
        cls._draw_arrowhead(draw, scaled_points[-2], scaled_points[-1], color, scale)

    @staticmethod
    def _edge_points(source: _NodeLayout, target: _NodeLayout) -> list[tuple[int, int]]:
        if source.lane_index == target.lane_index:
            source_center_x = source.x + source.width // 2
            target_center_x = target.x + target.width // 2
            if source.y <= target.y:
                start = (source_center_x, source.y + source.height)
                end = (target_center_x, target.y)
            else:
                start = (source_center_x, source.y)
                end = (target_center_x, target.y + target.height)
            mid_y = (start[1] + end[1]) // 2
            return [start, (start[0], mid_y), (end[0], mid_y), end]

        if source.lane_index < target.lane_index:
            start = (source.x + source.width, source.y + source.height // 2)
            end = (target.x, target.y + target.height // 2)
        else:
            start = (source.x, source.y + source.height // 2)
            end = (target.x + target.width, target.y + target.height // 2)
        mid_x = (start[0] + end[0]) // 2
        return [start, (mid_x, start[1]), (mid_x, end[1]), end]

    @staticmethod
    def _draw_arrowhead(
        draw: ImageDraw.ImageDraw,
        previous: tuple[int, int],
        end: tuple[int, int],
        color: str,
        scale: int,
    ) -> None:
        angle = math.atan2(end[1] - previous[1], end[0] - previous[0])
        size = 8 * scale
        spread = math.pi / 7
        points = [
            end,
            (
                int(end[0] - size * math.cos(angle - spread)),
                int(end[1] - size * math.sin(angle - spread)),
            ),
            (
                int(end[0] - size * math.cos(angle + spread)),
                int(end[1] - size * math.sin(angle + spread)),
            ),
        ]
        draw.polygon(points, fill=color)

    @staticmethod
    def _scale_box(box: tuple[int, int, int, int], scale: int) -> tuple[int, int, int, int]:
        return (box[0] * scale, box[1] * scale, box[2] * scale, box[3] * scale)

    @staticmethod
    def _line_height(font: _PillowFont) -> int:
        left, top, right, bottom = font.getbbox("Ag汉")
        return max(1, int(bottom - top + (right - left) // 100))

    @classmethod
    def _wrap_text(
        cls,
        draw: ImageDraw.ImageDraw,
        text: str,
        font: _PillowFont,
        max_width: int,
        max_lines: int,
    ) -> list[str]:
        if int(draw.textlength(text, font=font)) <= max_width:
            return [text]
        lines: list[str] = []
        current = ""
        for char in text:
            candidate = current + char
            if current and int(draw.textlength(candidate, font=font)) > max_width:
                lines.append(current)
                current = char.strip() if char == " " else char
                if len(lines) == max_lines:
                    return cls._ellipsize_lines(draw, lines, font, max_width)
            else:
                current = candidate
        if current:
            lines.append(current)
        if len(lines) > max_lines:
            lines = lines[:max_lines]
            return cls._ellipsize_lines(draw, lines, font, max_width)
        return lines

    @staticmethod
    def _ellipsize_lines(
        draw: ImageDraw.ImageDraw,
        lines: list[str],
        font: _PillowFont,
        max_width: int,
    ) -> list[str]:
        if not lines:
            return ["..."]
        last = lines[-1]
        suffix = "..."
        while last and int(draw.textlength(last + suffix, font=font)) > max_width:
            last = last[:-1]
        lines[-1] = f"{last}{suffix}" if last else suffix
        return lines

    @staticmethod
    def _load_font(size: int, bold: bool = False) -> _PillowFont:
        regular_candidates = [
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/Hiragino Sans GB.ttc",
            "/System/Library/Fonts/STHeiti Medium.ttc",
            "/Library/Fonts/Arial Unicode.ttf",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ]
        bold_candidates = [
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Medium.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
            "/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ]
        for font_path in bold_candidates if bold else regular_candidates:
            if not Path(font_path).is_file():
                continue
            try:
                return ImageFont.truetype(font_path, size=size)
            except OSError:
                continue
        return ImageFont.load_default(size=size)

    @staticmethod
    def _new_canvas(width: int, height: int, color: tuple[int, int, int]) -> list[bytearray]:
        row = bytearray(color * width)
        return [bytearray(row) for _ in range(height)]

    @staticmethod
    def _hex_rgb(value: str) -> tuple[int, int, int]:
        raw = value.strip().removeprefix("#")
        if len(raw) != 6:
            return (255, 255, 255)
        return (int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))

    @classmethod
    def _draw_rect(
        cls,
        canvas: list[bytearray],
        x: int,
        y: int,
        width: int,
        height: int,
        fill: tuple[int, int, int],
        stroke: tuple[int, int, int] | None,
    ) -> None:
        max_y = min(len(canvas), y + height)
        max_x = min(len(canvas[0]) // 3, x + width)
        for py in range(max(0, y), max_y):
            row = canvas[py]
            for px in range(max(0, x), max_x):
                offset = px * 3
                row[offset : offset + 3] = bytes(fill)
        if stroke is None:
            return
        cls._draw_line(canvas, x, y, x + width, y, stroke)
        cls._draw_line(canvas, x, y + height, x + width, y + height, stroke)
        cls._draw_line(canvas, x, y, x, y + height, stroke)
        cls._draw_line(canvas, x + width, y, x + width, y + height, stroke)

    @staticmethod
    def _draw_line(
        canvas: list[bytearray],
        x1: int,
        y1: int,
        x2: int,
        y2: int,
        color: tuple[int, int, int],
    ) -> None:
        dx = abs(x2 - x1)
        dy = -abs(y2 - y1)
        step_x = 1 if x1 < x2 else -1
        step_y = 1 if y1 < y2 else -1
        err = dx + dy
        x = x1
        y = y1
        while True:
            if 0 <= y < len(canvas) and 0 <= x < len(canvas[0]) // 3:
                offset = x * 3
                canvas[y][offset : offset + 3] = bytes(color)
            if x == x2 and y == y2:
                break
            err2 = 2 * err
            if err2 >= dy:
                err += dy
                x += step_x
            if err2 <= dx:
                err += dx
                y += step_y

    @classmethod
    def _draw_arrow(
        cls,
        canvas: list[bytearray],
        x1: int,
        y1: int,
        x2: int,
        y2: int,
        color: tuple[int, int, int],
    ) -> None:
        if abs(x2 - x1) >= abs(y2 - y1):
            direction = 1 if x2 >= x1 else -1
            cls._draw_line(canvas, x2, y2, x2 - 9 * direction, y2 - 5, color)
            cls._draw_line(canvas, x2, y2, x2 - 9 * direction, y2 + 5, color)
        else:
            direction = 1 if y2 >= y1 else -1
            cls._draw_line(canvas, x2, y2, x2 - 5, y2 - 9 * direction, color)
            cls._draw_line(canvas, x2, y2, x2 + 5, y2 - 9 * direction, color)

    @classmethod
    def _draw_text(
        cls,
        canvas: list[bytearray],
        x: int,
        y: int,
        text: str,
        color: tuple[int, int, int],
        scale: int = 1,
    ) -> None:
        cursor = x
        for char in text.upper()[:18]:
            glyph = _FONT_5X7.get(char, _FONT_5X7[" "])
            for row_idx, row in enumerate(glyph):
                for col_idx, value in enumerate(row):
                    if value != "1":
                        continue
                    for sy in range(scale):
                        py = y + row_idx * scale + sy
                        if py < 0 or py >= len(canvas):
                            continue
                        for sx in range(scale):
                            px = cursor + col_idx * scale + sx
                            if 0 <= px < len(canvas[0]) // 3:
                                offset = px * 3
                                canvas[py][offset : offset + 3] = bytes(color)
            cursor += 6 * scale

    @staticmethod
    def _preview_label(value: str, index: int) -> str:
        labels = {
            "页面结构": "LAYOUT",
            "核心交互": "ACTIONS",
            "输出区域": "OUTPUT",
            "状态反馈": "STATUS",
            "交互层": "UI",
            "Agent 编排层": "AGENT",
            "技能执行层": "SKILLS",
            "产物与质量层": "QUALITY",
            "质量层": "QUALITY",
            "左侧导航": "NAV",
            "Agent 列表": "AGENTS",
            "中间聊天区": "CHAT",
            "任务输入框": "INPUT",
            "技能快捷入口": "SKILLS",
            "ACP WS / SSE 切换": "SSE/WS",
            "右侧文件预览": "FILES",
            "校验结果": "VERIFY",
            "Spec 详情": "SPEC",
            "执行事件": "EVENTS",
            "下载操作": "DOWNLOAD",
            "错误提示": "ERRORS",
            "用户入口": "USER",
            "MiniMax 风格工作台": "WORKBENCH",
            "无状态 Agent Runtime": "RUNTIME",
            "JSON Agent 配置": "CONFIG",
            "结构化 Spec Builder": "SPEC",
            "Skill Registry": "REGISTRY",
            "Skill Runner": "RUNNER",
            "Artifact Store": "ARTIFACTS",
            "Preview Service": "PREVIEW",
            "内容级 Verifier": "VERIFIER",
            "自动修正重试": "RETRY",
            "Agent 最终回复": "REPLY",
        }
        return labels.get(value, f"N{index + 1}")

    @staticmethod
    def _encode_png(canvas: list[bytearray]) -> bytes:
        height = len(canvas)
        width = len(canvas[0]) // 3 if height else 0
        raw = b"".join(b"\x00" + bytes(row) for row in canvas)

        def chunk(kind: bytes, data: bytes) -> bytes:
            payload = kind + data
            return struct.pack(">I", len(data)) + payload + struct.pack(">I", zlib.crc32(payload) & 0xFFFFFFFF)

        ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
        return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b"")

    @classmethod
    def _build_xlsx(cls, rows: list[list[str]], summary_rows: list[list[str]]) -> bytes:
        files = {
            "[Content_Types].xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
</Types>""",
            "_rels/.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
            "xl/_rels/workbook.xml.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>""",
            "xl/workbook.xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="设备模板" sheetId="1" r:id="rId1"/>
    <sheet name="状态汇总" sheetId="2" r:id="rId2"/>
  </sheets>
</workbook>""",
            "xl/styles.xml": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="2"><font><sz val="11"/></font><font><b/><sz val="11"/></font></fonts>
  <fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFE2F0D9"/></patternFill></fill></fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/><xf numFmtId="0" fontId="1" fillId="1" borderId="0" applyFont="1" applyFill="1"/></cellXfs>
</styleSheet>""",
            "xl/worksheets/sheet1.xml": cls._sheet_xml(rows, freeze=True, validation=True),
            "xl/worksheets/sheet2.xml": cls._sheet_xml(summary_rows, freeze=False, validation=False),
        }
        return cls._zip_bytes(files)

    @staticmethod
    def _sheet_xml(rows: list[list[str]], freeze: bool, validation: bool) -> str:
        row_xml: list[str] = []
        for row_idx, row in enumerate(rows, start=1):
            cells: list[str] = []
            for col_idx, value in enumerate(row, start=1):
                cell_ref = f"{chr(64 + col_idx)}{row_idx}"
                style = ' s="1"' if row_idx == 1 else ""
                if value.startswith("="):
                    cells.append(f'<c r="{cell_ref}"{style}><f>{xml_escape(value[1:])}</f></c>')
                else:
                    cells.append(f'<c r="{cell_ref}" t="inlineStr"{style}><is><t>{xml_escape(value)}</t></is></c>')
            row_xml.append(f'<row r="{row_idx}">{"".join(cells)}</row>')
        pane = '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" state="frozen"/></sheetView></sheetViews>' if freeze else ""
        validations = (
            '<dataValidations count="1"><dataValidation type="list" allowBlank="1" sqref="D2:D100">'
            '<formula1>"在线,离线,告警"</formula1></dataValidation></dataValidations>'
            if validation
            else ""
        )
        return (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f"{pane}<sheetData>{''.join(row_xml)}</sheetData>{validations}</worksheet>"
        )

    @staticmethod
    def _pptx_bullets(kind: str, raw_bullets: object = None) -> list[str]:
        if isinstance(raw_bullets, list):
            bullets = []
            for item in raw_bullets:
                if item is None:
                    continue
                text = str(item).strip()
                if text:
                    bullets.append(text)
            if bullets:
                return bullets[:7]
        return {
            "cover": ["无状态运行", "结构化生成", "内容级校验"],
            "agenda": ["架构", "流程", "计划", "风险", "验收"],
            "architecture": ["JSON Agent 配置", "Skill Registry", "Artifact Store", "Verifier"],
            "process": ["自然语言", "结构化 Spec", "Skill 执行", "预览与校验", "自动重试"],
            "plan": ["事件流", "高质量模板", "Workbench 预览", "CLI/API 对齐"],
            "risks": ["无视觉证据不输出识别分数", "生成后必须解析校验", "失败时保留可追踪事件"],
            "summary": ["统一协议", "更诚实的检测", "更可验收的文件生成"],
        }.get(kind, ["结构化内容", "可预览产物", "可验证结果"])

    @classmethod
    def _build_pptx(cls, slides: list[dict[str, Any]]) -> bytes:
        if not slides:
            slides = [
                {
                    "title": "JetLinks Agent Runtime v2",
                    "kind": "cover",
                    "bullets": cls._pptx_bullets("cover"),
                }
            ]

        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)
        blank_layout = deck.slide_layouts[6]

        for idx, slide_spec in enumerate(slides, start=1):
            title = str(slide_spec.get("title") or f"Slide {idx}")
            kind = str(slide_spec.get("kind") or "content")
            raw_bullets = slide_spec.get("bullets")
            bullets = [str(item) for item in raw_bullets] if isinstance(raw_bullets, list) else cls._pptx_bullets(kind)
            slide = deck.slides.add_slide(blank_layout)
            cls._style_slide(slide, idx, len(slides), title, kind, bullets)

        buffer = io.BytesIO()
        deck.save(buffer)
        return buffer.getvalue()

    @classmethod
    def _style_slide(
        cls,
        slide: Any,
        index: int,
        total: int,
        title: str,
        kind: str,
        bullets: list[str],
    ) -> None:
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = cls._rgb(248, 250, 252)

        accent = cls._pptx_accent(index - 1)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        if kind == "cover" and index == 1:
            title_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.15), Inches(11.55), Inches(1.1))
            cls._set_text(title_box, title, size=34, bold=True, color=cls._rgb(15, 23, 42), align=PP_ALIGN.CENTER)
            subtitle = " · ".join(bullets[:3])
            subtitle_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.45), Inches(10.95), Inches(0.5))
            cls._set_text(subtitle_box, subtitle, size=17, color=cls._rgb(71, 85, 105), align=PP_ALIGN.CENTER)
        else:
            tag_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.35), Inches(4.2), Inches(0.35))
            cls._set_text(tag_box, cls._pptx_kind_label(kind), size=10, bold=True, color=accent)

            title_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.78), Inches(11.85), Inches(0.9))
            cls._set_text(title_box, title, size=28, bold=True, color=cls._rgb(15, 23, 42))

            body_box = slide.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(11.2), Inches(4.35))
            frame = body_box.text_frame
            frame.clear()
            frame.word_wrap = True
            for bullet_idx, bullet in enumerate(bullets[:7]):
                paragraph = frame.paragraphs[0] if bullet_idx == 0 else frame.add_paragraph()
                paragraph.text = f"- {bullet}"
                paragraph.font.name = "PingFang SC"
                paragraph.font.size = Pt(19)
                paragraph.font.color.rgb = cls._rgb(30, 41, 59)
                paragraph.space_after = Pt(10)

        footer_box = slide.shapes.add_textbox(Inches(0.72), Inches(6.95), Inches(11.85), Inches(0.28))
        cls._set_text(
            footer_box,
            f"{index:02d} / {total:02d}  JetLinks Agent Runtime v2",
            size=9,
            color=cls._rgb(100, 116, 139),
            align=PP_ALIGN.RIGHT,
        )

    @staticmethod
    def _set_text(
        shape: Any,
        text: str,
        *,
        size: int,
        color: RGBColor,
        bold: bool = False,
        align: PP_ALIGN | None = None,
    ) -> None:
        frame = shape.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        if align is not None:
            paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = "PingFang SC"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    @staticmethod
    def _pptx_accent(index: int) -> RGBColor:
        colors = [(37, 99, 235), (22, 163, 74), (217, 119, 6), (124, 58, 237), (220, 38, 38), (8, 145, 178)]
        return SkillRunner._rgb(*colors[index % len(colors)])

    @staticmethod
    def _rgb(red: int, green: int, blue: int) -> RGBColor:
        return RGBColor(red, green, blue)  # type: ignore[no-untyped-call]

    @staticmethod
    def _pptx_kind_label(kind: str) -> str:
        labels = {
            "agenda": "AGENDA",
            "architecture": "ARCHITECTURE",
            "process": "PROCESS",
            "plan": "PLAN",
            "risks": "RISKS",
            "summary": "SUMMARY",
        }
        return labels.get(kind, "CONTENT")

    @staticmethod
    def _json_text(value: dict[str, Any]) -> str:
        return json.dumps(value, ensure_ascii=False, indent=2)
