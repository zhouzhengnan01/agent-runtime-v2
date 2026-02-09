#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def _project_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _default_output_path() -> Path:
    out_dir = _project_root() / "storage" / "skill_outputs" / "pptx-generation"
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / "智巡云SaaS产品方向规划.pptx"


def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 定义颜色
    TITLE_COLOR = RGBColor(31, 78, 120)
    ACCENT_COLOR = RGBColor(68, 114, 196)
    TEXT_COLOR = RGBColor(64, 64, 64)

    # 封面页
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "智巡云 SaaS 产品方向规划"
    title_frame.paragraphs[0].font.size = Pt(54)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(0.5))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "AI赋能的企业级SaaS服务平台"
    subtitle_frame.paragraphs[0].font.size = Pt(24)
    subtitle_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
    subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # 目录页
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "目录"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "1. 产品定位与愿景"
    for item in ["2. 核心能力：API开放平台", "3. 业务方向一：微小企业AI工具箱",
                 "4. 业务方向二：巡检与复判服务", "5. 商业模式与市场策略",
                 "6. 技术架构与竞争优势", "7. 实施路线图"]:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(20)
        p.space_before = Pt(12)

    # 第1页：产品定位与愿景
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "1. 产品定位与愿景"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "核心定位"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    for item in ["AI赋能的企业级SaaS服务平台",
                 "为微小企业提供低成本、高效率的数字化解决方案",
                 "聚焦垂直行业的智能化升级"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)
        p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "目标客户"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    for item in ["微小型企业（50人以下）",
                 "传统行业数字化转型企业",
                 "需要智能巡检服务的企业"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)
        p.space_before = Pt(8)

    # 第2页：API开放平台
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "2. 核心能力：API开放平台"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "对外开放AI能力接口"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    for item in ["智能识别、数据分析、内容生成", "支持按调用次数计费"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "无缝集成已有系统"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    for item in ["明道云、企业微信等第三方平台", "Webhook、OAuth标准协议支持"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "开发者生态"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    for item in ["API文档、SDK、示例代码", "开发者社区与技术支持"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "参考：https://www.explinks.com/apihub"
    p.font.size = Pt(14)
    p.font.italic = True
    p.space_before = Pt(16)

    # 第3页：业务方向一 - 概览
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "3. 微小企业AI工具箱"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    modules = [
        ("AI营销获客系统", "智能推荐企业、线索挖掘"),
        ("智能邮件营销", "历史邮件学习、个性化生成"),
        ("AI内容运营", "小红书+微信+视频生成"),
        ("AI招投标助手", "文档解析、自动生成标书"),
        ("AI财务运营", "票据识别、报表自动生成"),
        ("AI市场舆情分析", "实时监控、品牌情感分析")
    ]

    for i, (name, desc) in enumerate(modules):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{name}"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR
        if i > 0:
            p.space_before = Pt(12)

        p = tf.add_paragraph()
        p.text = desc
        p.level = 1
        p.font.size = Pt(16)

    # 第4页：AI营销获客
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "AI营销获客系统"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "核心功能"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    for item in ["智能推荐合适企业客户",
                 "构建企业信息库，精准判断客户意向",
                 "集成谷歌搜索接口，自动化线索挖掘"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "价值主张"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    for item in ["降低获客成本 60%", "提升线索质量 3倍"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    # 第5页：AI内容运营
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "AI内容运营"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "全渠道覆盖"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    for item in ["小红书自动化运营", "微信公众号内容生成", "AI视频生成与剪辑"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "价值主张"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    for item in ["节约运营成本 70%", "快速获客，提升品牌曝光"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    # 第6页：其他AI工具
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "更多AI工具"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "AI招投标助手"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    p = tf.add_paragraph()
    p.text = "节约标书制作时间 80%，提升中标率"
    p.level = 1
    p.font.size = Pt(16)

    p = tf.add_paragraph()
    p.text = "AI财务运营"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    p = tf.add_paragraph()
    p.text = "票据智能识别、财务报表自动生成、税务合规检查"
    p.level = 1
    p.font.size = Pt(16)

    p = tf.add_paragraph()
    p.text = "AI市场舆情分析"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    p = tf.add_paragraph()
    p.text = "基于搜索引擎的实时监控、品牌情感分析、竞品追踪"
    p.level = 1
    p.font.size = Pt(16)

    # 第7页：巡检与复判 - 概览
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "4. 巡检与复判服务"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "三大核心能力"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    services = [
        ("无人机巡检SaaS平台", "对接大疆无人机，实现自动化巡检"),
        ("AR眼镜巡检方案", "对接雷鸟AR眼镜，第一视角巡检"),
        ("智能复判SaaS服务", "AI+人工混合复判，确保准确率")
    ]

    for name, desc in services:
        p = tf.add_paragraph()
        p.text = name
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR
        p.space_before = Pt(12)

        p = tf.add_paragraph()
        p.text = desc
        p.level = 1
        p.font.size = Pt(16)

    # 第8页：无人机巡检
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "无人机巡检SaaS平台"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "技术对接"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    p = tf.add_paragraph()
    p.text = "大疆无人机SDK集成"
    p.level = 1
    p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "https://developer.dji.com/cn/"
    p.level = 1
    p.font.size = Pt(14)
    p.font.italic = True

    p = tf.add_paragraph()
    p.text = "核心功能"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    for item in ["飞行任务规划与调度", "实时视频流监控", "巡检数据自动采集与分析"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    # 第9页：AR眼镜巡检
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "AR眼镜巡检方案"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "设备对接"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    p = tf.add_paragraph()
    p.text = "雷鸟AR眼镜SDK集成"
    p.level = 1
    p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "https://open.rayneo.cn/"
    p.level = 1
    p.font.size = Pt(14)
    p.font.italic = True

    p = tf.add_paragraph()
    p.text = "核心功能"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    for item in ["第一视角巡检记录", "AR辅助标注与指引", "语音交互与远程协助"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "优势：用户自主选择设备型号，多品牌兼容"
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    # 第10页：智能复判
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "智能复判SaaS服务"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "核心功能"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    for item in ["AI自动识别缺陷", "人工+AI混合复判机制", "接口对接已有系统"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "应用场景"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    p = tf.add_paragraph()
    p.text = "电力巡检、设备检测、质量把控"
    p.level = 1
    p.font.size = Pt(18)

    p = tf.add_paragraph()
    p.text = "价值主张"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)

    for item in ["提升准确率，降低漏检率", "历史数据沉淀，形成知识库"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)

    # 第11页：商业模式
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "5. 商业模式与市场策略"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "定价策略"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    pricing = [
        ("免费版", "个人/初创", "免费", "基础功能，有使用限制"),
        ("基础版", "微小企业", "¥299/月", "核心功能包"),
        ("专业版", "中小企业", "¥999/月", "全功能 + API调用"),
        ("企业版", "大型企业", "定制报价", "私有化 + 定制开发")
    ]

    for version, target, price, features in pricing:
        p = tf.add_paragraph()
        p.text = f"{version} - {price}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.space_before = Pt(8)

        p = tf.add_paragraph()
        p.text = f"目标：{target} | {features}"
        p.level = 1
        p.font.size = Pt(14)

    p = tf.add_paragraph()
    p.text = "盈利模式：订阅费 + API按量计费 + 增值服务 + 硬件捆绑"
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    # 第12页：技术架构
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "6. 技术架构与竞争优势"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "技术架构"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    for item in ["响应式Web + 移动端App",
                 "微服务架构，云原生部署",
                 "自研+第三方AI模型集成",
                 "等保三级认证，数据加密"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(16)

    p = tf.add_paragraph()
    p.text = "核心竞争力"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(16)

    for item in ["多设备兼容能力", "垂直行业数据沉淀",
                 "开放API生态", "低成本高效率", "易用性"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(16)

    # 第13页：实施路线图
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "7. 实施路线图"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    roadmap = [
        ("第一阶段（MVP - 3个月）", [
            "开发API开放平台基础能力",
            "上线AI营销获客 + 邮件营销模块",
            "完成大疆无人机SDK对接"
        ]),
        ("第二阶段（功能扩展 - 6个月）", [
            "上线内容运营、招投标、财务模块",
            "完成雷鸟AR眼镜对接",
            "建立智能复判服务体系"
        ]),
        ("第三阶段（生态建设 - 12个月）", [
            "完善开发者生态与API市场",
            "拓展行业解决方案",
            "建立渠道代理体系"
        ])
    ]

    for i, (stage, items) in enumerate(roadmap):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = stage
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR
        if i > 0:
            p.space_before = Pt(12)

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.level = 1
            p.font.size = Pt(14)

    # 总结页
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "总结"
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "核心价值主张"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "让微小企业也能享受AI带来的效率革命"
    p.font.size = Pt(24)
    p.font.color.rgb = TITLE_COLOR
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(16)

    p = tf.add_paragraph()
    p.text = "三大支柱"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(30)

    for item in ["开放API平台 - 技术赋能",
                 "AI工具箱 - 降本增效",
                 "巡检复判 - 垂直行业突破"]:
        p = tf.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(18)
        p.space_before = Pt(8)

    # Q&A页
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    qa_box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1.5))
    qa_frame = qa_box.text_frame
    qa_frame.text = "Q & A"
    qa_frame.paragraphs[0].font.size = Pt(72)
    qa_frame.paragraphs[0].font.bold = True
    qa_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    qa_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    thanks_box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(0.5))
    thanks_frame = thanks_box.text_frame
    thanks_frame.text = "感谢观看！"
    thanks_frame.paragraphs[0].font.size = Pt(24)
    thanks_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
    thanks_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    return prs

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="生成 PowerPoint 演示文稿（PPTX）")
    parser.add_argument(
        "--output",
        help="输出文件路径（.pptx）；不传则输出到 storage/skill_outputs/pptx-generation/",
    )
    args = parser.parse_args()

    prs = create_presentation()
    output_path = Path(args.output).expanduser().resolve() if args.output else _default_output_path()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"PPT已生成：{output_path}")
