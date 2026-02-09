#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path
from typing import Iterable, Optional


logger = logging.getLogger("deliverables-export")


def _read_text(value: Optional[str], *, encoding: str = "utf-8") -> str:
    if not value:
        return ""
    raw = value.strip()
    if raw.startswith("@"):
        path = Path(raw[1:]).expanduser().resolve()
        return path.read_text(encoding=encoding)
    path = Path(raw).expanduser()
    if path.exists() and path.is_file():
        return path.resolve().read_text(encoding=encoding)
    return value


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _write_utf8(path: Path, content: str) -> None:
    _ensure_parent(path)
    path.write_text(content, encoding="utf-8")


_RE_HEADING = re.compile(r"^\s*(#{1,6})\s+(.*)\s*$")
_RE_NUMBER = re.compile(r"^\s*(\d+)\s*[\.\:：]\s*(.*)\s*$")
_RE_BULLET = re.compile(r"^\s*[-*]\s+(.*)\s*$")


def _iter_markdown_blocks(lines: Iterable[str]):
    in_code = False
    code_lines: list[str] = []
    for raw in lines:
        line = raw.rstrip("\n")
        if line.strip().startswith("```"):
            if in_code:
                yield ("code", "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            yield ("blank", "")
            continue

        m = _RE_HEADING.match(line)
        if m:
            level = min(3, max(1, len(m.group(1))))
            yield (f"h{level}", m.group(2).strip())
            continue

        m = _RE_NUMBER.match(line)
        if m:
            yield ("number", m.group(2).strip())
            continue

        m = _RE_BULLET.match(line)
        if m:
            yield ("bullet", m.group(1).strip())
            continue

        yield ("text", line)

    if in_code and code_lines:
        yield ("code", "\n".join(code_lines))


def _build_docx_from_markdown(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as exc:
        raise RuntimeError("python-docx not available. Install it first (pip install python-docx).") from exc

    doc = Document()

    for kind, text in _iter_markdown_blocks(markdown.splitlines()):
        if kind.startswith("h"):
            level = int(kind[1:])
            doc.add_heading(text, level=level)
            continue
        if kind == "bullet":
            try:
                doc.add_paragraph(text, style="List Bullet")
            except Exception:
                doc.add_paragraph(f"- {text}")
            continue
        if kind == "number":
            try:
                doc.add_paragraph(text, style="List Number")
            except Exception:
                doc.add_paragraph(f"1. {text}")
            continue
        if kind == "code":
            para = doc.add_paragraph()
            run = para.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            continue
        if kind == "blank":
            doc.add_paragraph("")
            continue
        doc.add_paragraph(text)

    import io

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Export deliverables: commands txt + steps docx.")
    parser.add_argument("--commands-in", help="命令调用内容：文件路径或 '@file' 或直接文本")
    parser.add_argument("--steps-in", help="步骤说明内容：文件路径或 '@file' 或直接文本（支持简单 Markdown）")
    parser.add_argument("--out-dir", help="输出目录（默认生成 命令调用.txt + 步骤说明.docx）")
    parser.add_argument("--commands-out", help="命令文件输出路径（.txt）")
    parser.add_argument("--steps-out", help="步骤文件输出路径（.docx）")
    parser.add_argument("--commands-name", default="命令调用.txt", help="命令文件名（配合 --out-dir）")
    parser.add_argument("--steps-name", default="步骤说明.docx", help="步骤文件名（配合 --out-dir）")
    parser.add_argument("--log-level", default="INFO", help="日志级别（DEBUG/INFO/WARNING/ERROR）")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=getattr(logging, str(args.log_level).upper(), logging.INFO),
        format="%(asctime)s - %(levelname)s - %(message)s",
    )

    commands = _read_text(args.commands_in)
    steps = _read_text(args.steps_in)

    if not commands.strip():
        logger.error("commands 内容为空：请提供 --commands-in")
        return 2
    if not steps.strip():
        logger.error("steps 内容为空：请提供 --steps-in")
        return 2

    if args.out_dir:
        out_dir = Path(args.out_dir).expanduser().resolve()
        commands_out = out_dir / args.commands_name
        steps_out = out_dir / args.steps_name
    else:
        commands_out = Path(args.commands_out).expanduser().resolve() if args.commands_out else None
        steps_out = Path(args.steps_out).expanduser().resolve() if args.steps_out else None

    if not commands_out or not steps_out:
        logger.error("缺少输出路径：请提供 --out-dir，或同时提供 --commands-out 与 --steps-out")
        return 2

    logger.info("写入命令文件: %s", commands_out)
    _write_utf8(commands_out, commands)

    logger.info("生成 DOCX: %s", steps_out)
    docx_bytes = _build_docx_from_markdown(steps)
    _ensure_parent(steps_out)
    steps_out.write_bytes(docx_bytes)

    logger.info("完成：%s (%d chars), %s (%d bytes)", commands_out, len(commands), steps_out, len(docx_bytes))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

