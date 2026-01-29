"""
Daily review report API.
"""

from __future__ import annotations

import io
import os
import re
import time
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, HTTPException, Query
from pydantic import BaseModel, Field
from fastapi.responses import StreamingResponse

from app.services.review_report_service import get_review_report_service

router = APIRouter()


class ReviewReportGenerateRequest(BaseModel):
    date: Optional[str] = Field(default=None, description="Report date (YYYY-MM-DD)")
    base_url: Optional[str] = Field(default=None, description="Base URL")
    agent_id: Optional[str] = Field(default=None, description="Agent ID")
    sample_size: Optional[int] = Field(default=None, description="Sample size")
    force: bool = Field(default=False, description="Force regenerate")


def _escape_html(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _parse_markdown_blocks(markdown: str) -> List[Dict[str, Any]]:
    blocks: List[Dict[str, Any]] = []
    in_code = False
    code_lines: List[str] = []

    for raw_line in (markdown or "").splitlines():
        line = raw_line.rstrip("\n")
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                blocks.append({"type": "code", "text": "\n".join(code_lines)})
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            blocks.append({"type": "blank", "text": ""})
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            blocks.append({"type": "heading", "level": max(1, min(3, level)), "text": text})
            continue
        if re.match(r"^[-*+]\s+", stripped):
            text = re.sub(r"^[-*+]\s+", "", stripped)
            blocks.append({"type": "bullet", "text": text})
            continue
        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            blocks.append({"type": "number", "text": text})
            continue
        blocks.append({"type": "paragraph", "text": stripped})

    if code_lines:
        blocks.append({"type": "code", "text": "\n".join(code_lines)})

    return blocks


def _pick_pdf_font() -> str:
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception:
        return "Helvetica"

    candidates = [
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/arphic/ukai.ttc",
        "/usr/share/fonts/truetype/arphic/uming.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
    ]

    for path in candidates:
        if not os.path.exists(path):
            continue
        try:
            font_name = "ReportFont"
            if font_name in pdfmetrics.getRegisteredFontNames():
                return font_name
            if path.lower().endswith(".ttc"):
                pdfmetrics.registerFont(TTFont(font_name, path, subfontIndex=0))
            else:
                pdfmetrics.registerFont(TTFont(font_name, path))
            return font_name
        except Exception:
            continue

    return "Helvetica"


def _build_pdf(markdown: str, title: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"reportlab not available: {e}")

    font_name = _pick_pdf_font()
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "Base",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=14,
    )
    heading1 = ParagraphStyle("Heading1", parent=base, fontSize=16, leading=20, spaceBefore=6, spaceAfter=6)
    heading2 = ParagraphStyle("Heading2", parent=base, fontSize=13, leading=18, spaceBefore=6, spaceAfter=4)
    heading3 = ParagraphStyle("Heading3", parent=base, fontSize=11.5, leading=16, spaceBefore=4, spaceAfter=3)
    code_style = ParagraphStyle(
        "Code",
        parent=base,
        fontName="Courier",
        fontSize=9,
        leading=12,
        leftIndent=8,
        backColor="#f5f5f5",
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=title,
    )
    story = []

    blocks = _parse_markdown_blocks(markdown)
    for block in blocks:
        text = block.get("text") or ""
        if block.get("type") == "heading":
            level = int(block.get("level") or 1)
            style = heading1 if level == 1 else heading2 if level == 2 else heading3
            story.append(Paragraph(_escape_html(text), style))
            continue
        if block.get("type") == "bullet":
            story.append(Paragraph(f"- {_escape_html(text)}", base))
            continue
        if block.get("type") == "number":
            story.append(Paragraph(f"1. {_escape_html(text)}", base))
            continue
        if block.get("type") == "code":
            if text:
                html = _escape_html(text).replace("\n", "<br/>")
                story.append(Paragraph(html, code_style))
            else:
                story.append(Spacer(1, 4))
            continue
        if block.get("type") == "blank":
            story.append(Spacer(1, 6))
            continue
        if text:
            story.append(Paragraph(_escape_html(text), base))
        story.append(Spacer(1, 4))

    if not story:
        story.append(Paragraph("(empty report)", base))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


def _build_docx(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"python-docx not available: {e}")

    doc = Document()
    blocks = _parse_markdown_blocks(markdown)

    for block in blocks:
        text = block.get("text") or ""
        if block.get("type") == "heading":
            level = int(block.get("level") or 1)
            doc.add_heading(text, level=min(3, max(1, level)))
            continue
        if block.get("type") == "bullet":
            try:
                doc.add_paragraph(text, style="List Bullet")
            except Exception:
                doc.add_paragraph(f"- {text}")
            continue
        if block.get("type") == "number":
            try:
                doc.add_paragraph(text, style="List Number")
            except Exception:
                doc.add_paragraph(f"1. {text}")
            continue
        if block.get("type") == "code":
            para = doc.add_paragraph()
            run = para.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            continue
        if block.get("type") == "blank":
            doc.add_paragraph("")
            continue
        doc.add_paragraph(text)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


@router.get("/review-reports")
async def list_review_reports(
    pageIndex: int = Query(default=0, ge=0, description="Page index (0-based)"),
    pageSize: int = Query(default=50, ge=1, le=200, description="Page size"),
) -> Dict[str, Any]:
    store = get_review_report_service()
    result = store.list_reports(page_index=pageIndex, page_size=pageSize)
    return {
        "status": 200,
        "message": "success",
        "result": result,
        "timestamp": int(time.time() * 1000),
    }


@router.get("/review-reports/{report_date}")
async def get_review_report(report_date: str) -> Dict[str, Any]:
    store = get_review_report_service()
    report = store.get_report(report_date)
    if not report:
        raise HTTPException(status_code=404, detail="report not found")
    return {
        "status": 200,
        "message": "success",
        "result": report,
        "timestamp": int(time.time() * 1000),
    }


@router.post("/review-reports/generate")
async def generate_review_report(req: ReviewReportGenerateRequest) -> Dict[str, Any]:
    store = get_review_report_service()
    report_date = store._normalize_date(req.date)
    existing = store.get_report(report_date)
    if existing and not req.force:
        return {
            "status": 200,
            "message": "success",
            "result": {"report": existing, "generated": False},
            "timestamp": int(time.time() * 1000),
        }

    report = await store.generate_report(
        date_value=report_date,
        base_url=req.base_url,
        agent_id=req.agent_id,
        sample_size=req.sample_size,
        force=bool(req.force),
        generated_by="manual",
    )
    if not report:
        raise HTTPException(status_code=400, detail="no records for report")

    return {
        "status": 200,
        "message": "success",
        "result": {"report": report, "generated": True},
        "timestamp": int(time.time() * 1000),
    }


@router.get("/review-reports/{report_date}/export")
async def export_review_report(
    report_date: str,
    format: str = Query(default="pdf", description="pdf or docx"),
) -> StreamingResponse:
    store = get_review_report_service()
    report = store.get_report(report_date)
    if not report:
        raise HTTPException(status_code=404, detail="report not found")

    markdown = str(report.get("report_markdown") or "")
    if not markdown:
        markdown = "(empty report)"

    fmt = str(format or "pdf").lower().strip()
    if fmt not in {"pdf", "docx", "word"}:
        raise HTTPException(status_code=400, detail="format must be pdf or docx")

    filename_base = f"review-report-{report_date}"
    if fmt == "pdf":
        payload = _build_pdf(markdown, title=filename_base)
        media_type = "application/pdf"
        filename = f"{filename_base}.pdf"
    else:
        payload = _build_docx(markdown)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"{filename_base}.docx"

    headers = {"Content-Disposition": f"attachment; filename=\"{filename}\""}
    return StreamingResponse(io.BytesIO(payload), media_type=media_type, headers=headers)
