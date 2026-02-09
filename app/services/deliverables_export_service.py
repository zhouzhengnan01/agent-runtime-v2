from __future__ import annotations

import io
import logging
import os
import re
import time
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

logger = logging.getLogger(__name__)

_RE_HEADING = re.compile(r"^\s*(#{1,6})\s+(.*)\s*$")
_RE_NUMBER = re.compile(r"^\s*(\d+)\s*[\.\:：]\s*(.*)\s*$")
_RE_BULLET = re.compile(r"^\s*[-*]\s+(.*)\s*$")


def _now_ms() -> int:
    return int(time.time() * 1000)


def _project_root() -> Path:
    # Resolve repo root by locating main.py (repo entrypoint).
    current = Path(__file__).resolve()
    for parent in current.parents:
        if (parent / "main.py").exists() and (parent / "app").exists():
            return parent
    # Fallback (expected layout: <root>/app/services/<file>)
    return current.parents[2]


def _resolve_output_dir(raw: Optional[str]) -> Path:
    if not raw:
        return _project_root() / "storage" / "uploads" / "documents" / "deliverables"
    expanded = os.path.expanduser(raw)
    path = Path(expanded)
    if not path.is_absolute():
        path = _project_root() / path
    return path.resolve()


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _write_text_utf8(path: Path, content: str) -> None:
    _ensure_parent(path)
    path.write_text(content, encoding="utf-8")


def _iter_blocks(lines: Iterable[str]) -> Iterable[Tuple[str, str]]:
    in_code = False
    code_lines: List[str] = []

    for raw in lines:
        line = raw.rstrip("\n")
        if line.strip().startswith("```"):
            if in_code:
                yield ("code", "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            yield ("blank", "")
            continue

        m = _RE_HEADING.match(line)
        if m:
            level = min(3, max(1, len(m.group(1))))
            yield (f"h{level}", m.group(2).strip())
            continue

        m = _RE_NUMBER.match(line)
        if m:
            yield ("number", m.group(2).strip())
            continue

        m = _RE_BULLET.match(line)
        if m:
            yield ("bullet", m.group(1).strip())
            continue

        yield ("text", line)

    if in_code and code_lines:
        yield ("code", "\n".join(code_lines))


def _build_docx_from_markdown(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as exc:
        raise RuntimeError(f"python-docx not available: {exc}") from exc

    doc = Document()
    for kind, text in _iter_blocks((markdown or "").splitlines()):
        if kind.startswith("h"):
            level = int(kind[1:])
            doc.add_heading(text, level=min(3, max(1, level)))
            continue
        if kind == "bullet":
            try:
                doc.add_paragraph(text, style="List Bullet")
            except Exception:
                doc.add_paragraph(f"- {text}")
            continue
        if kind == "number":
            try:
                doc.add_paragraph(text, style="List Number")
            except Exception:
                doc.add_paragraph(f"1. {text}")
            continue
        if kind == "code":
            para = doc.add_paragraph()
            run = para.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            continue
        if kind == "blank":
            doc.add_paragraph("")
            continue
        doc.add_paragraph(text)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _build_static_url(path: Path) -> Optional[str]:
    """Map a filesystem path under storage/uploads -> /static/uploads/ URL."""
    try:
        storage_dir = _project_root() / "storage" / "uploads"
        rel = path.resolve().relative_to(storage_dir.resolve())
        return f"/static/uploads/{rel.as_posix()}"
    except Exception:
        return None


@dataclass(frozen=True)
class DeliverablesExportResult:
    job_id: str
    created_at_ms: int
    commands: Dict[str, Any]
    steps: Dict[str, Any]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "jobId": self.job_id,
            "createdAt": self.created_at_ms,
            "commands": self.commands,
            "steps": self.steps,
        }


def export_deliverables(
    *,
    commands_text: str,
    steps_markdown: str,
    output_dir: Optional[str] = None,
    commands_filename: str = "命令调用.txt",
    steps_filename: str = "步骤说明.docx",
    commands_out_path: Optional[str] = None,
    steps_out_path: Optional[str] = None,
) -> DeliverablesExportResult:
    if not (commands_text or "").strip():
        raise ValueError("commands_text is required")
    if not (steps_markdown or "").strip():
        raise ValueError("steps_markdown is required")

    job_id = uuid.uuid4().hex[:10]
    created_at = _now_ms()

    base_dir = _resolve_output_dir(output_dir)
    default_dir = base_dir / job_id

    if commands_out_path:
        commands_path = Path(os.path.expanduser(commands_out_path)).resolve()
    else:
        commands_path = default_dir / commands_filename

    if steps_out_path:
        steps_path = Path(os.path.expanduser(steps_out_path)).resolve()
    else:
        steps_path = default_dir / steps_filename

    logger.info(
        "Deliverables export start job=%s dir=%s commands_chars=%d steps_chars=%d",
        job_id,
        str(default_dir),
        len(commands_text),
        len(steps_markdown),
    )

    _write_text_utf8(commands_path, commands_text)
    commands_size = commands_path.stat().st_size
    logger.info("Deliverables export wrote commands: %s (%d bytes)", str(commands_path), commands_size)

    steps_bytes = _build_docx_from_markdown(steps_markdown)
    _ensure_parent(steps_path)
    steps_path.write_bytes(steps_bytes)
    logger.info("Deliverables export wrote steps: %s (%d bytes)", str(steps_path), len(steps_bytes))

    commands_url = _build_static_url(commands_path)
    steps_url = _build_static_url(steps_path)

    return DeliverablesExportResult(
        job_id=job_id,
        created_at_ms=created_at,
        commands={
            "path": str(commands_path),
            "url": commands_url,
            "size": commands_size,
            "filename": commands_path.name,
        },
        steps={
            "path": str(steps_path),
            "url": steps_url,
            "size": len(steps_bytes),
            "filename": steps_path.name,
        },
    )
