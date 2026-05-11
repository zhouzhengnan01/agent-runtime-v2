from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation

from app.core.artifacts import ArtifactStore
from app.core.skills.runner_types import SkillRunResult
from app.core.skills.runner import SkillRunner
from app.core.workflow.verifier import OutputVerifier


def test_pptx_generation_outputs_openable_office_deck(tmp_path: Path) -> None:
    store = ArtifactStore(root_dir=tmp_path)
    paths = store.prepare_thread("pptx-openable")
    result = SkillRunner(store).run(
        "pptx-generation",
        {
            "title": "JetLinks IoT 平台方案",
            "slides": [
                {"title": "JetLinks IoT 平台方案", "kind": "cover"},
                {"title": "目录", "kind": "agenda"},
                {"title": "总体架构", "kind": "architecture", "bullets": ["自定义要点", "设备接入", "规则引擎"]},
                {"title": "处理流程", "kind": "process"},
                {"title": "实施计划", "kind": "plan"},
                {"title": "风险控制", "kind": "risks"},
                {"title": "总结", "kind": "summary"},
            ],
        },
        paths,
    )

    artifact = result.outputs[0]
    deck_path = store.resolve_virtual_path(artifact.thread_id, artifact.path)

    with zipfile.ZipFile(deck_path) as archive:
        names = set(archive.namelist())

    assert artifact.name == "deck.pptx"
    assert artifact.kind == "presentation"
    assert "ppt/presentation.xml" in names
    assert "ppt/theme/theme1.xml" in names
    assert "ppt/slideMasters/slideMaster1.xml" in names
    assert any(name.startswith("ppt/slideLayouts/") for name in names)
    assert all(f"ppt/slides/_rels/slide{idx}.xml.rels" in names for idx in range(1, 8))
    assert {"docProps/core.xml", "docProps/app.xml"} <= names

    deck = Presentation(str(deck_path))
    assert len(deck.slides) == 7
    text = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "JetLinks IoT 平台方案" in text
    assert "自定义要点" in text

    verification = OutputVerifier(store).verify({"skill_name": "pptx-generation"}, result)
    assert verification.passed, verification.failed_checks


def test_coco_verifier_fails_when_auto_annotation_has_no_detected_objects(tmp_path: Path) -> None:
    store = ArtifactStore(root_dir=tmp_path)
    paths = store.prepare_thread("empty-coco")
    artifact = store.write_text_artifact(
        paths,
        "annotations.coco.json",
        '{"images":[{"id":1,"file_name":"scene.jpg","width":640,"height":480}],"annotations":[],"categories":[{"id":1,"name":"person"}]}',
    )

    verification = OutputVerifier(store).verify(
        {"skill_name": "data-auto-annotation"},
        SkillRunResult(skill_name="data-auto-annotation", outputs=[artifact]),
    )

    assert not verification.passed
    assert "coco_has_detected_objects" in verification.failed_checks
